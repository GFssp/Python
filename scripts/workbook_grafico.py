from openpyxl.chart.axis import DateAxis
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from datetime import date
from time import sleep

# criando a planilha
wb = Workbook()

# acessando a aba
ws = wb.active

# criando os dados
rows = [
    ['Date', 'Batch 1', 'Batch 2', 'Batch 3'],
    [date(2015,9, 1), 40, 30, 25],
    [date(2015,9, 2), 40, 25, 30],
    [date(2015,9, 3), 50, 30, 45],
    [date(2015,9, 4), 30, 25, 40],
    [date(2015,9, 5), 25, 35, 30],
    [date(2015,9, 6), 20, 40, 35],
]

# inserindo os dados
for row in rows:
    ws.append(row)

# criando o gráfico
c1 = LineChart()
c1.title = "Line Chart"
c1.style = 13
c1.y_axis.title = 'Tamanho'
c1.x_axis.title = 'Número de testes'

data = Reference(ws, min_col=2, min_row=1, max_col=4, max_row=7)

# adicionando as referências
c1.add_data(data, titles_from_data=True)

# adicionando o gráfico à planilha
ws.add_chart(c1, "A10")

for value in ws.iter_rows(values_only=True):
  print(value)
sleep(3)

# salvando a planilha
wb.save(filename='line_chart.xlsx')
